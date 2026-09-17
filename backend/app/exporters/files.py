from io import BytesIO, StringIO
from zipfile import ZipFile, ZIP_DEFLATED
import pandas as pd
from openpyxl.styles import Font

SHEETS=[("store","Store"),("website_data","Website_Data"),("categories","Categories"),("products","Products"),("products_flat","Products_Flat"),("product_categories","Product_Categories"),("images","Images"),("product_options","Product_Options"),("option_values","Option_Values"),("variants","Variants"),("variant_option_values","Variant_Option_Values"),("tags","Tags"),("product_tags","Product_Tags"),("seo","SEO"),("validation_issues","Validation_Issues")]
def readme_lines(session):
    store=(session.tables.get("store") or [{}])[0]
    product_limit="N/A" if session.data_type!="PRODUCTS" else (session.max_products or "ALL")
    lines=["Salla Store Migration Preparation Export",f"Export timestamp: {store.get('Extraction_Date','')}",f"Store URL: {store.get('Store_URL','')}",f"Store name: {store.get('Store_Name','')}",f"Data mode: {session.mode}",f"Data type: {session.data_type}",f"Product limit: {product_limit}","Extractor version: 2.0.0","", "Blank cells may indicate data that was not publicly available from the source storefront."]
    if session.mode=="MOCK":lines += ["","WARNING: This export contains MOCK DATA generated for Preview testing and must not be treated as extracted store data."]
    lines += ["", "Sheets:"]+[f"- {name}: normalized {key.replace('_',' ')} records" for key,name in SHEETS]
    return lines
def xlsx_bytes(session):
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        pd.DataFrame({"README":readme_lines(session)}).to_excel(writer,sheet_name="README",index=False)
        for key,name in SHEETS:
            rows=session.tables.get(key,[]); pd.DataFrame(rows).to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws["A1"].font=Font(bold=True)
            for col in ws.columns:
                letter=col[0].column_letter; ws.column_dimensions[letter].width=min(55,max(12,max((len(str(c.value or "")) for c in col),default=10)+2))
    return out.getvalue()
def zip_bytes(session):
    out=BytesIO()
    with ZipFile(out,"w",ZIP_DEFLATED) as z:
        z.writestr("README.txt","\ufeff"+"\n".join(readme_lines(session)))
        for key,_ in SHEETS:
            csv=pd.DataFrame(session.tables.get(key,[])).to_csv(index=False)
            z.writestr(f"{key}.csv","\ufeff"+csv)
    return out.getvalue()
