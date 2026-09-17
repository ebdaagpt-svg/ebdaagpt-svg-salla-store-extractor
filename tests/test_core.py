import io, time, zipfile
from tempfile import TemporaryDirectory
import pytest
from openpyxl import load_workbook
from backend.app.utils import normalize_url, validate_public_host, number_or_none, bool_or_none, html_to_text, URLSafetyError
from backend.app.mock.data import mock_catalog
from backend.app.models import Session
from backend.app.validators.catalog import validate_catalog
from backend.app.exporters.files import xlsx_bytes, zip_bytes
from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.extractors.salla import ExtractionFailure, embedded_json_objects, extract_size_volume, extraction_scope, parse_sitemap, product_entries_from_dom, rate_limit_delay, source_id_from_url, product_from_page, json_objects, website_data_from_page
from backend.app.services.session_store import SessionStore
from backend.app.transformers.products_flat import build_products_flat
from bs4 import BeautifulSoup

def session():
    tables=mock_catalog();issues=validate_catalog(tables);tables["validation_issues"]=[x.model_dump() for x in issues]
    return Session(id="test",stage="DEMO_MODE",mode="MOCK",message="test",tables=tables,issues=issues)
def test_url_validation(): assert normalize_url("https://shop.example.com")=="https://shop.example.com/"
def test_complex_salla_filter_url():
    value=normalize_url("https://shop.example.com/products?filters[category_id]=12&filters[available]=1")
    assert "filters%5Bcategory_id%5D=12" in value and "filters%5Bavailable%5D=1" in value
@pytest.mark.parametrize("url",["file:///etc/passwd","ftp://example.com","http://localhost:8000"])
def test_bad_urls(url):
    with pytest.raises(URLSafetyError): normalize_url(url)
def test_ssrf_ip():
    with pytest.raises(URLSafetyError): validate_public_host("http://127.0.0.1/")
def test_normalizers():
    assert number_or_none("SAR 12.50")==12.5 and number_or_none("") is None
    assert bool_or_none("متوفر") is True and bool_or_none("unknown") is None
    assert html_to_text("<p>مرحبا <b>بك</b></p>")=="مرحبا بك"
def test_sitemap_parsing_and_product_ids():
    xml='''<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><sitemap><loc>https://shop.example/sitemap-1.xml</loc></sitemap></sitemapindex>'''
    nested,rows=parse_sitemap(xml); assert nested==["https://shop.example/sitemap-1.xml"] and rows==[]
    xml='''<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9" xmlns:image="http://www.google.com/schemas/sitemap-image/1.1"><url><loc>https://shop.example/item/p123</loc><lastmod>2026-01-01</lastmod><image:image><image:loc>https://cdn.example/a.jpg</image:loc></image:image></url></urlset>'''
    nested,rows=parse_sitemap(xml); assert not nested and rows[0]["images"]==["https://cdn.example/a.jpg"]
    assert source_id_from_url(rows[0]["url"])=="123"
def test_exact_url_scope_and_category_dom_isolation():
    assert extraction_scope("https://shop.example/")=="STORE"
    assert extraction_scope("https://shop.example/item/p123")=="PRODUCT"
    assert extraction_scope("https://shop.example/ar/cables/c44")=="CATEGORY"
    assert extraction_scope("https://shop.example/products?filters[category_id]=44")=="CATEGORY"
    soup=BeautifulSoup('<a href="/unrelated/p999">outside</a><salla-product-card><a href="/inside/p123">inside</a></salla-product-card>',"lxml")
    assert [row["url"] for row in product_entries_from_dom(soup,"https://shop.example/category/c44")]==["https://shop.example/inside/p123"]
def test_jsonld_product_prices_and_public_fields():
    html='''<html><head><script type="application/ld+json">{"@type":"Product","productID":"123","name":"عبوة اختبار","description":"<p>وصف</p>","weight":{"value":500,"unitText":"ml"},"notes":"يحفظ مبرداً","offers":{"price":80,"highPrice":100,"priceCurrency":"SAR","availability":"https://schema.org/InStock","shippingDetails":{"shippingLabel":"جاهز للشحن"}}}</script></head></html>'''
    soup=BeautifulSoup(html,"lxml")
    product,*_=product_from_page(json_objects(soup),soup,"https://shop.example/item/p123")
    assert product["Price"]==100 and product["Sale_Price"]==80
    assert product["Size_Volume"]=="500.0 ml" and product["Shipping_Status"]=="جاهز للشحن"
def test_size_volume_regex():
    assert extract_size_volume("عطر مركز 100 مل")=="100 مل"
    assert extract_size_volume("Bottle 80ML")=="80 ML"
def test_rate_limit_backoff_and_retry_after():
    assert rate_limit_delay({},0)==3
    assert rate_limit_delay({},1)==4.5
    assert rate_limit_delay({},2)==5
    assert rate_limit_delay({"retry-after":"9"},0)==9
def test_public_salla_datalayer_categories():
    html='''<html><head><script type="application/ld+json">{"@type":"Product","productID":"123","name":"منتج","offers":{"price":10}}</script></head><body><script>window.dataLayer.push({"event":"detail","ecommerce":{"detail":{"products":[{"id":123,"name":"منتج","categories":[{"id":44,"name":"العناية"}]}]}}});</script></body></html>'''
    soup=BeautifulSoup(html,"lxml")
    parsed=product_from_page(embedded_json_objects(soup),soup,"https://shop.example/item/p123")
    assert parsed and parsed[-1]==[("العناية","https://shop.example/c44")]
def test_website_data_normalization():
    html='''<html lang="ar"><head><link rel="canonical" href="https://shop.example/"><script type="application/ld+json">{"@type":"Organization","name":"متجر الاختبار","description":"وصف عام","logo":"/logo.png","telephone":"+966500000000","email":"hello@shop.example","address":{"streetAddress":"الشارع الأول","addressLocality":"الرياض","addressCountry":"SA"},"geo":{"latitude":24.7,"longitude":46.7},"sameAs":["https://instagram.com/shop"]}</script></head><body><a href="https://x.com/shop">X</a></body></html>'''
    soup=BeautifulSoup(html,"lxml")
    store,rows=website_data_from_page(embedded_json_objects(soup),soup,"https://shop.example/")
    values={row["Field"]:row["Value"] for row in rows}
    assert store[0]["Store_Name"]=="متجر الاختبار"
    assert values["Location_City"]=="الرياض" and values["Latitude"]=="24.7"
    assert values["Social_Instagram"]=="https://instagram.com/shop"
    assert values["Social_X_Twitter"]=="https://x.com/shop"
def test_salla_public_state_website_fields():
    html='''<html lang="ar"><head><script>salla.event.dispatchEvents({"twilight::init":{"store":{"id":77,"name":"متجر","username":"shop.user","country":"SA","contacts":{"mobile":"+966500000001","email":"info@shop.test","whatsapp":"+966500000002"},"social":{"instagram":"https://instagram.com/shop"},"settings":{"tax":{"number":"VAT123"},"commercial_number":"CR456"},"apps":{"appstore":"https://apps.apple.com/app/id1"}},"currencies":{"SAR":{"code":"SAR"}}}});</script></head></html>'''
    soup=BeautifulSoup(html,"lxml")
    raw=embedded_json_objects(soup)
    store,rows=website_data_from_page(raw,soup,"https://shop.test/")
    pairs={(row["Field"],row["Value"]) for row in rows}
    assert store[0]["Store_ID"]=="77" and store[0]["Currency"]=="SAR"
    assert ("Phone","+966500000001") in pairs
    assert ("Email","info@shop.test") in pairs
    assert ("VAT_ID","VAT123") in pairs and ("Commercial_Registration","CR456") in pairs
    assert ("Social_Instagram","https://instagram.com/shop") in pairs
def test_products_flat_one_row_per_product():
    tables=mock_catalog()
    rows=build_products_flat(tables)
    assert len(rows)==len(tables["products"])==20
    assert rows[0]["Category_Names"] and rows[0]["Main_Image_URL"]
    assert rows[0]["Image_Count"]==2
def test_relationships_and_demo_warnings():
    issues=validate_catalog(mock_catalog()); assert any(x.Field=="SKU" for x in issues) and any(x.Field=="Images" for x in issues)
def test_excel_generation():
    wb=load_workbook(io.BytesIO(xlsx_bytes(session()))); assert {"README","Website_Data","Products","Products_Flat","Validation_Issues"}.issubset(wb.sheetnames); assert wb["Store"]["G2"].value=="MOCK"
def test_csv_generation():
    z=zipfile.ZipFile(io.BytesIO(zip_bytes(session()))); assert "products.csv" in z.namelist() and "products_flat.csv" in z.namelist(); assert "MOCK DATA" in z.read("README.txt").decode("utf-8-sig")

def test_sqlite_session_and_export_persistence():
    with TemporaryDirectory() as directory:
        path=f"{directory}/sessions.sqlite3"
        first=SessionStore(path); current=session(); first.save(current); first.save_export(current.id,"csv",b"PK-test")
        restarted=SessionStore(path)
        assert restarted.get(current.id).stage=="DEMO_MODE"
        assert restarted.get_export(current.id,"csv")==b"PK-test"
        active=Session(id="active",stage="FETCHING",message="working"); restarted.save(active)
        restarted.save_checkpoint("active",1,{"product":{"id":"1"}})
        assert restarted.checkpoint_count("active")==1
        assert restarted.load_checkpoints("active")[0]["product"]["id"]=="1"
        assert restarted.recover_interrupted()==1
        assert restarted.get("active").stage=="ERROR"
        restarted.clear_checkpoints("active"); assert restarted.checkpoint_count("active")==0

def test_demo_mode_end_to_end(monkeypatch):
    import backend.app.main as main
    async def network_blocked(url,*args,**kwargs): raise ExtractionFailure("PREVIEW_NETWORK_RESTRICTED","blocked",True)
    monkeypatch.setattr(main,"validate_public_host",lambda url:url)
    monkeypatch.setattr(main,"extract_live",network_blocked)
    client=TestClient(app)
    assert client.get("/api/health").status_code==200
    result=client.post("/api/extract",json={"store_url":"https://example.salla.sa/","extraction_mode":"FULL"})
    assert result.status_code==202
    sid=result.json()["id"]
    for _ in range(50):
        status=client.get(f"/api/extraction/{sid}").json()
        if status["stage"] in {"READY","COMPLETED","PARTIAL_SUCCESS","DEMO_MODE","ERROR"}: break
        time.sleep(.01)
    assert status["mode"]=="MOCK" and status["extraction_mode"]=="FULL"
    tables=client.get(f"/api/extraction/{sid}/tables").json()["tables"]
    assert len(tables["products"])==20 and tables["store"][0]["Data_Mode"]=="MOCK"
    assert client.get(f"/api/extraction/{sid}/export/xlsx").content[:2]==b"PK"
    assert client.get(f"/api/extraction/{sid}/export/csv").content[:2]==b"PK"

def test_website_demo_mode_end_to_end(monkeypatch):
    import backend.app.main as main
    async def network_blocked(url,*args,**kwargs): raise ExtractionFailure("PREVIEW_NETWORK_RESTRICTED","blocked",True)
    monkeypatch.setattr(main,"validate_public_host",lambda url:url)
    monkeypatch.setattr(main,"extract_website_data",network_blocked)
    client=TestClient(app)
    result=client.post("/api/extract",json={"store_url":"https://example.salla.sa/","data_type":"WEBSITE"})
    assert result.status_code==202
    sid=result.json()["id"]
    for _ in range(50):
        status=client.get(f"/api/extraction/{sid}").json()
        if status["stage"] in {"READY","DEMO_MODE","ERROR"}: break
        time.sleep(.01)
    assert status["mode"]=="MOCK" and status["data_type"]=="WEBSITE"
    payload=client.get(f"/api/extraction/{sid}/tables").json()
    assert payload["tables"]["products"]==[]
    assert len(payload["tables"]["website_data"])>=5

def test_product_limit_validation_and_persistence(monkeypatch):
    import backend.app.main as main
    async def network_blocked(url,*args,**kwargs): raise ExtractionFailure("PREVIEW_NETWORK_RESTRICTED","blocked",True)
    monkeypatch.setattr(main,"validate_public_host",lambda url:url)
    monkeypatch.setattr(main,"extract_live",network_blocked)
    client=TestClient(app)
    assert client.post("/api/extract",json={"store_url":"https://example.salla.sa/","max_products":1001}).status_code==422
    result=client.post("/api/extract",json={"store_url":"https://example.salla.sa/","data_type":"PRODUCTS","max_products":50})
    sid=result.json()["id"]
    for _ in range(50):
        status=client.get(f"/api/extraction/{sid}").json()
        if status["stage"] in {"READY","DEMO_MODE","ERROR"}: break
        time.sleep(.01)
    assert status["max_products"]==50 and status["stats"]["product_limit"]==50
